"""Kobo API access, cleaning, consolidation and Excel export for the FFS seedling dashboard."""
import io
import random
from datetime import datetime, timedelta

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from species import SPECIES, WARDS

QTY_COLS = list(SPECIES.keys())
LABELS = {c: v["label"] for c, v in SPECIES.items()}
CATEGORY = {c: v["category"] for c, v in SPECIES.items()}
EXCEL_FONT = "Maiandra GD"


# ---------------------------------------------------------------- Kobo API
def fetch_submissions(server: str, token: str, asset_uid: str) -> list:
    url = f"{server.rstrip('/')}/api/v2/assets/{asset_uid}/data/"
    params = {"format": "json", "limit": 1000}
    headers = {"Authorization": f"Token {token}"}
    rows = []
    while url:
        r = requests.get(url, params=params, headers=headers, timeout=60)
        r.raise_for_status()
        j = r.json()
        rows.extend(j.get("results", []))
        url = j.get("next")
        params = None
    return rows


def flatten(rows: list) -> pd.DataFrame:
    """Kobo prefixes grouped fields (g_facilitator/ward). Keep the last path segment."""
    out = []
    for r in rows:
        d = {}
        for k, v in r.items():
            if isinstance(v, (list, dict)):
                continue
            d[k.rsplit("/", 1)[-1]] = v
        out.append(d)
    return pd.DataFrame(out)


# ---------------------------------------------------------------- cleaning
def prepare(df: pd.DataFrame, dedupe: bool = True):
    """Return (clean_df, info). info holds counts for the data-checks tab."""
    info = {"raw": len(df), "duplicates_removed": 0, "no_ward": 0, "under_3": 0}
    if df.empty:
        return df, info
    df = df.copy()
    for c in QTY_COLS:
        # Kobo only sends columns for species someone actually chose, so many are missing early on
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).astype(int) if c in df else 0
    for key in ("ind", "fruit"):
        for k in (1, 2, 3):
            qn, nn = f"other_{key}_{k}_qty", f"other_{key}_{k}_name"
            df[qn] = pd.to_numeric(df[qn], errors="coerce").fillna(0).astype(int) if qn in df else 0
            df[nn] = (df[nn].fillna("").astype(str).str.replace(r"\s+", " ", regex=True).str.strip() if nn in df else "")
    for c in ("subcounty_name", "ward_name", "ffs_group", "farmer_name", "farmer_phone",
              "facilitator_name", "physical_location", "aez_final"):
        if c not in df:
            df[c] = ""
        df[c] = df[c].fillna("").astype(str).str.strip()
    # fall back to the ward lookup if the calculated names are missing
    ward_code = df["ward"] if "ward" in df else pd.Series("", index=df.index)
    df["ward_name"] = df["ward_name"].where(df["ward_name"] != "", ward_code.map(lambda w: WARDS.get(w, {}).get("ward", "")))
    df["subcounty_name"] = df["subcounty_name"].where(df["subcounty_name"] != "", ward_code.map(lambda w: WARDS.get(w, {}).get("subcounty", "")))
    info["no_ward"] = int((df["ward_name"] == "").sum())
    df = df[df["ward_name"] != ""]

    df["ffs_group_clean"] = df["ffs_group"].str.replace(r"\s+", " ", regex=True).str.upper()
    blank = pd.Series([None] * len(df), index=df.index)
    df["farmer_no"] = pd.to_numeric(df["farmer_no"] if "farmer_no" in df else blank, errors="coerce")
    df["submitted"] = pd.to_datetime(df["_submission_time"] if "_submission_time" in df else blank, errors="coerce", utc=True)

    ind = [c for c in QTY_COLS if CATEGORY[c] == "Indigenous"]
    fru = [c for c in QTY_COLS if CATEGORY[c] != "Indigenous"]
    named_ind = sum((df[f"other_ind_{k}_name"] != "").astype(int) for k in (1, 2, 3))
    named_fru = sum((df[f"other_fruit_{k}_name"] != "").astype(int) for k in (1, 2, 3))
    df["total_other_ind"] = df[[f"other_ind_{k}_qty" for k in (1, 2, 3)]].sum(axis=1)
    df["total_other_fruit"] = df[[f"other_fruit_{k}_qty" for k in (1, 2, 3)]].sum(axis=1)
    df["n_indigenous"] = (df[ind] > 0).sum(axis=1) + named_ind
    df["n_fruit"] = (df[fru] > 0).sum(axis=1) + named_fru
    df["total_indigenous"] = df[ind].sum(axis=1) + df["total_other_ind"]
    df["total_fruit"] = df[fru].sum(axis=1) + df["total_other_fruit"]
    df["total_seedlings"] = df["total_indigenous"] + df["total_fruit"]
    info["under_3"] = int(((df["n_indigenous"] < 3) | (df["n_fruit"] < 3)).sum())

    if dedupe:
        before = len(df)
        df = (df.sort_values("submitted")
                .drop_duplicates(["ward_name", "ffs_group_clean", "farmer_no"], keep="last"))
        info["duplicates_removed"] = before - len(df)
    return df.reset_index(drop=True), info


# ---------------------------------------------------------------- consolidation
def _used_cols(df):
    return [c for c in QTY_COLS if c in df and df[c].sum() > 0]


OTHER_COLS = ["Sub-county", "Ward", "FFS group", "Category", "Species", "Seedlings"]


def others_long(df: pd.DataFrame) -> pd.DataFrame:
    """One row per 'other' species entry. Names are matched ignoring capitals and extra spaces."""
    parts = []
    for key, cat in (("ind", "Indigenous"), ("fruit", "Fruit / agroforestry")):
        for k in (1, 2, 3):
            n, q = f"other_{key}_{k}_name", f"other_{key}_{k}_qty"
            if n not in df or df.empty:
                continue
            sub = df[df[n] != ""]
            if not sub.empty:
                parts.append(pd.DataFrame({"Sub-county": sub["subcounty_name"], "Ward": sub["ward_name"],
                                           "FFS group": sub["ffs_group_clean"], "Category": cat,
                                           "key": sub[n].str.lower(), "raw": sub[n], "Seedlings": sub[q]}))
    if not parts:
        return pd.DataFrame(columns=OTHER_COLS)
    out = pd.concat(parts, ignore_index=True)
    shown = out.groupby("key")["raw"].agg(lambda x: x.value_counts().index[0])  # most common spelling
    shown = shown.map(lambda t: t[:1].upper() + t[1:])
    out["Species"] = out["key"].map(shown)
    return out[OTHER_COLS]


def others_by_group(df: pd.DataFrame) -> pd.DataFrame:
    o = others_long(df)
    if o.empty:
        return o
    g = o.groupby(["Sub-county", "Ward", "FFS group", "Category", "Species"], as_index=False).agg(
        Farmers=("Seedlings", "size"), Seedlings=("Seedlings", "sum"))
    return g


def county_by_species(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for c in QTY_COLS:
        rows.append({"Category": CATEGORY[c], "Species": LABELS[c],
                     "Farmers choosing": int((df[c] > 0).sum()) if len(df) else 0,
                     "Seedlings": int(df[c].sum()) if len(df) else 0})
    o = others_long(df) if len(df) else pd.DataFrame(columns=OTHER_COLS)
    if not o.empty:
        g = o.groupby(["Category", "Species"], as_index=False).agg(Farmers=("Seedlings", "size"), Seedlings=("Seedlings", "sum"))
        for _, r in g.iterrows():
            rows.append({"Category": r["Category"], "Species": f"{r['Species']} (other – not on list)",
                         "Farmers choosing": int(r["Farmers"]), "Seedlings": int(r["Seedlings"])})
    out = pd.DataFrame(rows)
    return out.sort_values(["Category", "Seedlings"], ascending=[True, False]).reset_index(drop=True)


def consolidate(df: pd.DataFrame, by: list, group_label: dict = None) -> pd.DataFrame:
    """One row per `by` combination, one column per species (only species with seedlings)."""
    if df.empty:
        return pd.DataFrame()
    cols = _used_cols(df)
    g = df.groupby(by, dropna=False)
    out = g[cols].sum()
    out.insert(0, "Farmers", g.size())
    oth = g["total_other_ind"].sum() + g["total_other_fruit"].sum()
    out["Total seedlings"] = out[cols].sum(axis=1) + oth
    out = out.rename(columns={c: LABELS[c] for c in cols})
    if oth.sum() > 0:
        out.insert(out.columns.get_loc("Total seedlings"), "Other species (see Other_species)", oth)
    out = out.reset_index()
    if group_label:
        out = out.rename(columns=group_label)
    return out


def group_table(df):
    return consolidate(df, ["subcounty_name", "ward_name", "ffs_group_clean"],
                       {"subcounty_name": "Sub-county", "ward_name": "Ward", "ffs_group_clean": "FFS group"})


def ward_table(df):
    return consolidate(df, ["subcounty_name", "ward_name"], {"subcounty_name": "Sub-county", "ward_name": "Ward"})


def subcounty_table(df):
    return consolidate(df, ["subcounty_name"], {"subcounty_name": "Sub-county"})


def farmer_table(df):
    cols = ["subcounty_name", "ward_name", "ffs_group", "farmer_no", "farmer_name", "farmer_phone",
            "physical_location", "facilitator_name", "aez_final", "submitted"] + _used_cols(df) + ["total_seedlings"]
    out = df[[c for c in cols if c in df]].copy()
    out["submitted"] = out["submitted"].dt.tz_convert("Africa/Nairobi").dt.strftime("%Y-%m-%d %H:%M")
    for key, label in (("ind", "Other indigenous (specified)"), ("fruit", "Other fruit (specified)")):
        out[label] = df.apply(lambda r: "; ".join(f"{r[f'other_{key}_{k}_name']} ({r[f'other_{key}_{k}_qty']})"
                                                  for k in (1, 2, 3) if r[f"other_{key}_{k}_name"] != ""), axis=1)
    out = out.rename(columns={**LABELS, "subcounty_name": "Sub-county", "ward_name": "Ward", "ffs_group": "FFS group",
                              "farmer_no": "Farmer no.", "farmer_name": "Farmer name", "farmer_phone": "Farmer phone",
                              "physical_location": "Location", "facilitator_name": "Facilitator", "aez_final": "AEZ",
                              "submitted": "Submitted", "total_seedlings": "Total seedlings"})
    return out


# ---------------------------------------------------------------- Excel
def to_excel(df: pd.DataFrame, include_farmers: bool) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        county_by_species(df).to_excel(xw, sheet_name="County_by_species", index=False)
        for name, t in (("FFS_groups", group_table(df)), ("Wards", ward_table(df)), ("Sub_counties", subcounty_table(df))):
            (t if not t.empty else pd.DataFrame({"note": ["No data yet"]})).to_excel(xw, sheet_name=name, index=False)
        ob = others_by_group(df) if not df.empty else pd.DataFrame()
        if not ob.empty:
            ob.to_excel(xw, sheet_name="Other_species", index=False)
        if include_farmers and not df.empty:
            farmer_table(df).to_excel(xw, sheet_name="Farmers", index=False)
    buf.seek(0)
    wb = load_workbook(buf)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=EXCEL_FONT, bold=(c.row == 1))
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="D9EAF7")
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[1:] or col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 40)
        ws.row_dimensions[1].height = 60
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- demo data
def demo_rows(n_farmers: int = 160, seed: int = 7) -> list:
    """Kobo-style JSON (grouped keys) so the demo exercises the same code path as live data."""
    rnd = random.Random(seed)
    ward_codes = list(WARDS)
    rows, uid = [], 1
    now = datetime.utcnow()
    groups = {w: [f"{WARDS[w]['ward'].split('/')[0].split()[0]} FFS {i}" for i in range(1, rnd.randint(2, 4))] for w in ward_codes}
    for _ in range(n_farmers):
        w = rnd.choice(ward_codes)
        info = WARDS[w]
        zone = info["aez_main"]
        if info["aez_alt"] and rnd.random() < 0.4:
            zone = info["aez_alt"]
        ok = [c for c, s in SPECIES.items() if zone in s["zones"]]
        ind = [c for c in ok if SPECIES[c]["category"] == "Indigenous"]
        fru = [c for c in ok if SPECIES[c]["category"] != "Indigenous"]
        pick = rnd.sample(ind, rnd.randint(3, min(6, len(ind)))) + rnd.sample(fru, rnd.randint(3, min(5, len(fru))))
        g = rnd.choice(groups[w])
        r = {"_id": uid, "_submission_time": (now - timedelta(hours=rnd.randint(0, 72))).strftime("%Y-%m-%dT%H:%M:%S"),
             "g_facilitator/facilitator_name": f"Facilitator {w[:3].title()}", "g_facilitator/ward": w,
             "g_facilitator/ffs_group": g if rnd.random() > 0.15 else g.lower() + " ",
             "g_farmer/farmer_no": str(rnd.randint(1, 30)), "g_farmer/farmer_name": f"Farmer {uid}",
             "g_farmer/farmer_phone": f"07{rnd.randint(10000000, 99999999)}", "g_farmer/physical_location": "Village",
             "g_zone/subcounty_name": info["subcounty"], "g_zone/ward_name": info["ward"], "g_zone/aez_final": zone}
        for c in pick:
            r[f"g_qty_ind/{c}" if SPECIES[c]["category"] == "Indigenous" else f"g_qty_fruit/{c}"] = str(rnd.choice([2, 3, 5, 5, 10, 10, 20]))
        if rnd.random() < 0.25:
            r["g_indigenous/other_ind_1_name"] = rnd.choice(["Grevillea robusta", "grevillea robusta ", "Cassia spectabilis", "Eucalyptus"])
            r["g_indigenous/other_ind_1_qty"] = str(rnd.choice([5, 10, 20]))
        if rnd.random() < 0.15:
            r["g_fruit/other_fruit_1_name"] = rnd.choice(["Guava", "guava", "Passion fruit", "Pawpaw"])
            r["g_fruit/other_fruit_1_qty"] = str(rnd.choice([3, 5, 10]))
        rows.append(r)
        uid += 1
    return rows
